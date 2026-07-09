#!/usr/bin/env python3
"""
One-time migration ETL: clean the legacy Google-Sheets/Excel export into a
normalized, typed JSON dataset that seeds the platform (Cosmos DB later).

Input : Business Development - Client Segmentation.xlsx
Output: src/data/dataset.json      (normalized domain data)
        src/data/data-quality.json (issues found + fixes applied)

Run:    python3 scripts/etl.py
"""
from __future__ import annotations
import json, math, datetime as dt, re, unicodedata
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "Business Development - Client Segmentation.xlsx"
OUT_DIR = ROOT / "src" / "data"
OUT_DIR.mkdir(parents=True, exist_ok=True)

TODAY = dt.date(2026, 7, 9)  # snapshot date of the export

# --- canonical industry taxonomy -------------------------------------------
INDUSTRY_CANON = {
    "financial": "Finance", "finance": "Finance",
    "fmcg": "FMCG",
    "fashion": "Fashion",
    "tech/telecom": "Tech/Telecom", "tech": "Tech/Telecom",
    "other": "Other",
    "fair": "Fair",
    "automotive": "Automotive",
    "consultancy": "Professional Services",
    "professional services": "Professional Services",
    "consulting/professional services": "Professional Services",
}

# operative-name -> brand-data-name (resolves the join mismatches)
ALIAS = {
    "Poste IT": "Poste Italiane",
    "SMBC/Clifford Chance": "SMBC",
    "DELL EMEA": "DELL Technologies",
    "Bridge partners": "Bridge Partners",
    "Generali - Taverna": "Generali",
}

dq: list[dict] = []


def flag(entity, key, issue, fix):
    dq.append({"entity": entity, "key": key, "issue": issue, "fix": fix})


def norm_industry(v):
    if v is None:
        return None
    key = str(v).strip().lower()
    return INDUSTRY_CANON.get(key, str(v).strip())


def slug(name: str) -> str:
    s = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-zA-Z0-9]+", "-", s).strip("-").lower()
    return s or "item"


def iso(v, entity="", key=""):
    """Coerce a cell to an ISO date string, flagging bad values."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    if s.upper() in ("N/A", "NA", "-", "TBD"):
        flag(entity, key, f"non-date text {s!r} in date field", "set null")
        return None
    try:
        return dt.datetime.strptime(s, "%Y-%m-%d").date().isoformat()
    except ValueError:
        flag(entity, key, f"unparseable date {s!r}", "set null")
        return None


def clean_str(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


wb = openpyxl.load_workbook(XLSX, data_only=True)

# =========================================================================
# 1) BRANDS  (Brands Operative pipeline  +  Brand Data scores)
# =========================================================================
op = wb["Brands Operative"]
op_rows = []
for r in range(2, op.max_row + 1):
    name = clean_str(op.cell(r, 1).value)
    if not name:
        continue
    poc = op.cell(r, 5).value
    if isinstance(poc, (int, float)):
        flag(name, "POC", f"numeric POC {poc!r}", "set null")
        poc = None
    op_rows.append({
        "name": name,
        "status": clean_str(op.cell(r, 2).value),
        "priority": clean_str(op.cell(r, 3).value),
        "owner": clean_str(op.cell(r, 4).value),
        "poc": clean_str(poc),
        "industry": norm_industry(op.cell(r, 6).value),
        "industryRaw": clean_str(op.cell(r, 6).value),
        "initialContact": iso(op.cell(r, 7).value, name, "initialContact"),
        "lastContact": iso(op.cell(r, 8).value, name, "lastContact"),
        "followUp": iso(op.cell(r, 9).value, name, "followUp"),
        "closingFailed": iso(op.cell(r, 10).value, name, "closingFailed"),
        "notes": clean_str(op.cell(r, 11).value),
    })

# Brand Data scores keyed by data-name
bd = wb["Brand Data"]
scores = {}
for r in range(2, 47):
    nm = clean_str(bd.cell(r, 1).value)
    if not nm:
        continue
    def num(c):
        v = bd.cell(r, c).value
        return float(v) if isinstance(v, (int, float)) else None
    scores[nm] = {
        "tempoMonths": num(4),
        "tempoScore": num(5),
        "closing": clean_str(bd.cell(r, 6).value),
        "process": clean_str(bd.cell(r, 7).value),
        "dealsClosed": num(8),
        "budget": num(9),
        "assumption": clean_str(bd.cell(r, 10).value),
        "budgetScore": num(11),
        "customizationScore": num(12),   # 5 low effort .. 1 high effort
        "accessibilityRaw": clean_str(bd.cell(r, 13).value),
        "accessibilityScore": num(14),   # 5 low .. 1 high
        "receptivityScore": num(15),     # 1 hard .. 5 easy
        "alignmentScore": num(16),
        "industry": norm_industry(bd.cell(r, 2).value),
    }

# merge scores onto operative leads via alias map
brands = []
matched = set()
for row in op_rows:
    data_key = ALIAS.get(row["name"], row["name"])
    sc = scores.get(data_key)
    if sc:
        matched.add(data_key)
    b = dict(row)
    b["id"] = slug(row["name"])
    b["aliases"] = [data_key] if data_key != row["name"] else []
    b["scored"] = sc is not None
    if sc:
        econ = [sc["budgetScore"], sc["customizationScore"], sc["tempoScore"]]
        ease = [sc["accessibilityScore"], sc["alignmentScore"], sc["receptivityScore"]]
        sc = dict(sc)
        sc["economicalEfficiency"] = round(sum(econ) / 3, 4) if all(x is not None for x in econ) else None
        sc["easeOfAccess"] = round(sum(ease) / 3, 4) if all(x is not None for x in ease) else None
        b["scores"] = sc
    # date sanity
    ic, lc = b["initialContact"], b["lastContact"]
    if ic and lc and lc < ic:
        flag(row["name"], "dates", f"lastContact {lc} < initialContact {ic}", "kept raw, flagged")
    if ic and ic > TODAY.isoformat():
        flag(row["name"], "initialContact", f"future date {ic}", "kept raw, flagged")
    brands.append(b)

for k in set(scores) - matched:
    flag(k, "join", "scored brand not matched to a pipeline lead", "kept as data-only (check alias)")

# =========================================================================
# 2) INDUSTRIES  (recompute the Brand Analysis scorecard from clean data)
# =========================================================================
COMPANIES_EU = {"Fair":180,"Fashion":120,"Finance":250,"FMCG":200,"Other":600,
                "Professional Services":80,"Tech/Telecom":220}
MUSIC_VIDEO = {"Fair":4,"Fashion":3,"Finance":3,"FMCG":2,"Other":2,
               "Professional Services":2,"Tech/Telecom":2}
VALUATION = {"Fair":"Medium","Fashion":"Medium","Finance":"Medium","FMCG":"Low",
             "Other":"Low","Professional Services":"Medium","Tech/Telecom":"High"}

by_ind: dict[str, list] = {}
for sc in scores.values():
    by_ind.setdefault(sc["industry"], []).append(sc)

def avg(vals):
    vals = [v for v in vals if v is not None]
    return round(sum(vals) / len(vals), 4) if vals else None

industries = []
for ind, rows in sorted(by_ind.items()):
    econ = avg([avg([s["budgetScore"], s["customizationScore"], s["tempoScore"]]) for s in rows])
    ease = avg([avg([s["accessibilityScore"], s["alignmentScore"], s["receptivityScore"]]) for s in rows])
    opened = len(rows)
    total = COMPANIES_EU.get(ind)
    industries.append({
        "name": ind,
        "opened": opened,
        "companiesEU": total,
        "approachedMarket": round(opened / total, 4) if total else None,
        "economicalEfficiency": econ,
        "easeOfAccess": ease,
        "avgBudget": round(sum(s["budget"] for s in rows if s["budget"] is not None) / opened) if opened else None,
        "musicVideoFit": MUSIC_VIDEO.get(ind),
        "valuation": VALUATION.get(ind),
    })

# =========================================================================
# 3) SALES STRATEGY playbook (per industry, qualitative)
# =========================================================================
ss = wb["Sales Strategy"]
STRAT_COLS = {2:"Finance",3:"Professional Services",4:"Tech/Telecom",5:"Fashion",
              6:"FMCG",7:"Fair",8:"Other"}
strat_rows = {3:"marketingNeed",4:"productNeed",5:"staffNeed",6:"valueProposition",7:"resistance"}
playbook = {v: {"industry": v} for v in STRAT_COLS.values()}
for rr, field in strat_rows.items():
    for cc, ind in STRAT_COLS.items():
        playbook[ind][field] = clean_str(ss.cell(rr, cc).value)

# action plan by receptivity tier (rows 14-18, cols B-E)
tiers = {2:"High",3:"Medium",4:"Resistant",5:"To open"}
action_labels = {15:"research",16:"productExp",17:"salesProcess",18:"focus"}
playbook_by_tier = {t: {"tier": t} for t in tiers.values()}
for cc, tier in tiers.items():
    for rr, lbl in action_labels.items():
        val = ss.cell(rr, cc).value
        playbook_by_tier[tier][lbl] = round(val, 3) if isinstance(val, (int, float)) else clean_str(val)

# closed/restricted sector market sizing (cols G-I, rows 15-21)
market_sizing = []
for rr in range(15, 22):
    sec = clean_str(ss.cell(rr, 7).value)
    if sec:
        market_sizing.append({
            "sector": sec,
            "companies": ss.cell(rr, 8).value,
            "marketSizeUsdBn": ss.cell(rr, 9).value,
        })

# =========================================================================
# 4) AGENTS / AGENCIES pipeline
# =========================================================================
ag = wb["AgentsAgencies"]
agents = []
for r in range(2, ag.max_row + 1):
    name = clean_str(ag.cell(r, 1).value)
    if not name:
        continue
    agents.append({
        "id": slug(name), "name": name,
        "status": clean_str(ag.cell(r, 2).value),
        "priority": clean_str(ag.cell(r, 3).value),
        "owner": clean_str(ag.cell(r, 4).value),
        "poc": clean_str(ag.cell(r, 5).value),
        "initialContact": iso(ag.cell(r, 6).value, name, "initialContact"),
        "lastContact": iso(ag.cell(r, 7).value, name, "lastContact"),
        "followUp": iso(ag.cell(r, 8).value, name, "followUp"),
        "notes": clean_str(ag.cell(r, 9).value),
    })

# =========================================================================
# assemble
# =========================================================================
dataset = {
    "meta": {
        "source": XLSX.name,
        "snapshotDate": TODAY.isoformat(),
        "generatedBy": "scripts/etl.py",
        "company": "OOVIE Studios",
        "owners": sorted({b["owner"] for b in brands if b["owner"]}),
        "industries": sorted(set(INDUSTRY_CANON.values())),
        "counts": {"brands": len(brands), "scored": sum(b["scored"] for b in brands),
                    "agents": len(agents), "industries": len(industries)},
    },
    "brands": brands,
    "industries": industries,
    "playbook": list(playbook.values()),
    "playbookByTier": list(playbook_by_tier.values()),
    "marketSizing": market_sizing,
    "agents": agents,
}

(OUT_DIR / "dataset.json").write_text(json.dumps(dataset, indent=2, ensure_ascii=False))
(OUT_DIR / "data-quality.json").write_text(json.dumps({"issues": dq, "count": len(dq)}, indent=2, ensure_ascii=False))

print(f"brands={len(brands)} scored={sum(b['scored'] for b in brands)} "
      f"agents={len(agents)} industries={len(industries)} dq_issues={len(dq)}")
print("wrote", OUT_DIR / "dataset.json")
